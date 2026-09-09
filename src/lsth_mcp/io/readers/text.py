"""Đọc file chữ thuần và file Word.

python-docx là phụ thuộc mềm: thiếu thì báo cách cài, không làm hỏng cả gói.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, List

from ...core.errors import UnsupportedFormatError
from ...core.provenance import Source
from ..base import ReadResult, Reader

ENCODINGS = ("utf-8-sig", "utf-8", "cp1258", "latin-1")


class TextReader(Reader):
    suffixes = (".txt", ".md", ".log")

    def read(self, path: Path, **_: Any) -> ReadResult:
        path = Path(path)
        for encoding in ENCODINGS:
            try:
                text = path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
            return ReadResult(path=path, text=text, sources=[Source(file=str(path))],
                              meta={"encoding": encoding, "chars": len(text)})
        raise UnsupportedFormatError(
            f"Không đoán được bảng mã của {path.name}.",
            hint="Mở bằng Notepad++ hoặc VS Code, lưu lại dạng UTF-8 rồi chạy lại.",
        )


class DocxReader(Reader):
    suffixes = (".docx",)

    def read(self, path: Path, **_: Any) -> ReadResult:
        try:
            import docx  # type: ignore
        except ImportError as exc:
            raise UnsupportedFormatError(
                "Chưa cài python-docx nên không đọc được file Word.",
                hint="Chạy: pip install python-docx",
            ) from exc

        path = Path(path)
        document = docx.Document(str(path))
        paragraphs: List[str] = [p.text for p in document.paragraphs]
        tables: List[List[List[str]]] = [
            [[cell.text for cell in row.cells] for row in table.rows]
            for table in document.tables
        ]
        return ReadResult(
            path=path,
            text="\n".join(paragraphs),
            data={"paragraphs": paragraphs, "tables": tables},
            sources=[Source(file=str(path))],
            meta={"paragraph_count": len(paragraphs), "table_count": len(tables)},
        )


def read_text(path: Path, **kwargs: Any) -> ReadResult:
    return TextReader().read(Path(path), **kwargs)


def read_docx(path: Path, **kwargs: Any) -> ReadResult:
    return DocxReader().read(Path(path), **kwargs)
