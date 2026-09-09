"""Kiểm thử lớp đọc/ghi — phần chức năng chính của gói này."""
from __future__ import annotations

import json
import unittest
from pathlib import Path

from lsth_mcp.core.errors import (
    PathOutsideRootError,
    UnsupportedFormatError,
    ValidationError,
    WriteNotAllowedError,
)
from lsth_mcp.io.readers.excel import read_excel
from lsth_mcp.io.readers.jsonio import read_json
from lsth_mcp.io.readers.tabular import read_csv
from lsth_mcp.io.registry import list_files, read_any, reader_for
from lsth_mcp.io.writers.excel import write_rows
from lsth_mcp.io.writers.safe import SafeWriter
from lsth_mcp.io.writers.structured import append_jsonl, write_csv, write_json, write_jsonl
from helpers import temp_settings, unlock

ROWS = [
    {"Style": "66P866", "Material Code": "FB-K0007-HD", "Consumption": 0.216},
    {"Style": "66P866", "Material Code": "TR-0021", "Consumption": 1.5},
]


def writer_for(settings, server="test_server", **kw):
    return SafeWriter.for_server(server, settings=settings, **kw)


class TestReaders(unittest.TestCase):
    def setUp(self):
        self.settings = temp_settings()
        unlock(self.settings, "test_server")

    def test_csv_roundtrip_keeps_line_numbers_as_source(self):
        writer = writer_for(self.settings)
        target = self.settings.work_dir / "bom.csv"
        write_csv(writer, target, ROWS)

        result = read_csv(target)
        self.assertEqual(len(result.rows), 2)
        self.assertEqual(result.rows[0]["Style"], "66P866")
        self.assertEqual(result.sources[0].row, 2)   # dòng 1 là tiêu đề

    def test_excel_roundtrip_and_header_detection(self):
        writer = writer_for(self.settings)
        target = self.settings.work_dir / "bom.xlsx"
        write_rows(writer, target, ROWS, sheet_name="BOM")

        result = read_excel(target, sheet="BOM")
        self.assertEqual(result.meta["sheet"], "BOM")
        # Dòng 1 là dấu bản nháp, nên tiêu đề phải được dò ra ở dòng 2.
        self.assertEqual(result.meta["header_row"], 2)
        self.assertEqual(len(result.rows), 2)
        self.assertEqual(result.rows[1]["Material Code"], "TR-0021")
        self.assertEqual(result.sources[0].sheet, "BOM")

    def test_excel_missing_sheet_lists_what_exists(self):
        writer = writer_for(self.settings)
        target = self.settings.work_dir / "b.xlsx"
        write_rows(writer, target, ROWS, sheet_name="BOM")
        with self.assertRaises(Exception) as ctx:
            read_excel(target, sheet="KHONG-CO")
        self.assertIn("BOM", str(ctx.exception))

    def test_jsonl_bad_line_names_the_line_number(self):
        path = self.settings.work_dir / "cases.jsonl"
        path.write_text('{"a": 1}\nkhong-phai-json\n', encoding="utf-8")
        with self.assertRaises(ValidationError) as ctx:
            read_json(path)
        self.assertIn("Dòng 2", str(ctx.exception))

        lenient = read_json(path, strict=False)
        self.assertEqual(len(lenient.rows), 1)
        self.assertEqual(len(lenient.warnings), 1)

    def test_registry_rejects_unknown_format_with_a_list_of_known_ones(self):
        with self.assertRaises(UnsupportedFormatError) as ctx:
            reader_for(Path("a.dwg"))
        self.assertIn(".xlsx", str(ctx.exception))

    def test_read_any_refuses_paths_outside_the_data_roots(self):
        with self.assertRaises(PathOutsideRootError):
            read_any("/etc/passwd", settings=self.settings)

    def test_list_files_marks_which_ones_are_readable(self):
        writer = writer_for(self.settings)
        write_json(writer, self.settings.work_dir / "a.json", {"x": 1})
        (self.settings.work_dir / "b.dwg").write_bytes(b"0")
        found = {f["name"]: f["readable"] for f in list_files(settings=self.settings)}
        self.assertTrue(found["a.json"])
        self.assertFalse(found["b.dwg"])


class TestSafeWriter(unittest.TestCase):
    def setUp(self):
        self.settings = temp_settings()

    def test_write_blocked_until_gate_opens(self):
        writer = writer_for(self.settings, "s2_erp")
        with self.assertRaises(WriteNotAllowedError) as ctx:
            writer.write_text(self.settings.work_dir / "x.txt", "nội dung")
        self.assertIn("5/5", str(ctx.exception))   # chưa mã nào chạy song song

        unlock(self.settings, "s2_erp")
        result = writer_for(self.settings, "s2_erp").write_text(
            self.settings.work_dir / "x.txt", "nội dung")
        self.assertTrue(result.path.exists())

    def test_global_read_only_switch_beats_an_open_gate(self):
        settings = temp_settings(write_enabled=False)
        unlock(settings, "s2_erp")
        with self.assertRaises(WriteNotAllowedError) as ctx:
            writer_for(settings, "s2_erp").write_text(settings.work_dir / "x.txt", "a")
        self.assertIn("chỉ đọc", str(ctx.exception))

    def test_cannot_write_into_raw(self):
        unlock(self.settings, "test_server")
        with self.assertRaises(PathOutsideRootError):
            writer_for(self.settings).write_text(self.settings.raw_dir / "goc.txt", "a")

    def test_overwrite_keeps_a_backup(self):
        unlock(self.settings, "test_server")
        target = self.settings.work_dir / "tldg.txt"
        writer_for(self.settings).write_text(target, "bản 1")
        result = writer_for(self.settings).write_text(target, "bản 2")

        self.assertFalse(result.created)
        self.assertIsNotNone(result.backup_path)
        self.assertEqual(result.backup_path.read_text(encoding="utf-8"), "bản 1")
        self.assertEqual(target.read_text(encoding="utf-8"), "bản 2")

    def test_executables_are_refused(self):
        unlock(self.settings, "test_server")
        with self.assertRaises(ValidationError):
            writer_for(self.settings).write_text(self.settings.work_dir / "x.sh", "rm -rf /")

    def test_dry_run_touches_nothing(self):
        unlock(self.settings, "test_server")
        target = self.settings.work_dir / "thu.txt"
        writer_for(self.settings, dry_run=True).write_text(target, "a")
        self.assertFalse(target.exists())

    def test_failed_write_leaves_no_partial_file(self):
        unlock(self.settings, "test_server")
        target = self.settings.work_dir / "hong.json"

        def explode(_tmp: Path) -> None:
            raise RuntimeError("hỏng giữa chừng")

        with self.assertRaises(RuntimeError):
            writer_for(self.settings).write_with(target, explode)
        self.assertFalse(target.exists())
        self.assertEqual(list(self.settings.work_dir.glob(".tmp-*")), [])


class TestStructuredWriters(unittest.TestCase):
    def setUp(self):
        self.settings = temp_settings()
        unlock(self.settings, "test_server")
        self.writer = writer_for(self.settings)

    def test_json_writes_dataclasses_via_to_dict(self):
        from lsth_mcp.core.models import BomLine
        line = BomLine(style="66P866", material_code="FB-K0007-HD", consumption=0.216)
        target = self.settings.out_dir / "bom.json"
        write_json(self.writer, target, {"lines": [line]})
        data = json.loads(target.read_text(encoding="utf-8"))
        self.assertEqual(data["lines"][0]["style"], "66P866")

    def test_jsonl_append_keeps_previous_lines(self):
        target = self.settings.work_dir / "log.jsonl"
        write_jsonl(self.writer, target, [{"a": 1}])
        append_jsonl(writer_for(self.settings), target, {"a": 2})
        rows = read_json(target).rows
        self.assertEqual([r["a"] for r in rows], [1, 2])

    def test_csv_uses_bom_so_excel_shows_vietnamese(self):
        target = self.settings.out_dir / "v.csv"
        write_csv(self.writer, target, [{"tên": "Tài liệu đóng gói"}])
        self.assertTrue(target.read_bytes().startswith(b"\xef\xbb\xbf"))

    def test_draft_banner_and_needs_human_highlight(self):
        from openpyxl import load_workbook
        target = self.settings.out_dir / "nhap.xlsx"
        write_rows(self.writer, target, ROWS, needs_human={0: ["Consumption"]})
        ws = load_workbook(target).active
        self.assertIn("BẢN NHÁP", ws.cell(row=1, column=1).value)
        marked = ws.cell(row=3, column=3)          # dòng dữ liệu đầu, cột Consumption
        self.assertEqual(marked.fill.fgColor.rgb[-6:], "FFF2CC")


if __name__ == "__main__":
    unittest.main()
